import types

import pytest
import requests

from stylometry import ingest

HTML = ("<html><head><style>x{}</style></head><body><nav>menu</nav><script>j()</script>"
        "<p>Hello world.</p><footer>foot</footer></body></html>")


def test_text_and_markdown(tmp_path):
    for name in ("a.txt", "a.md"):
        (tmp_path / name).write_text("héllo", encoding="utf-8")
        assert ingest.load_text(str(tmp_path / name)) == "héllo"


def test_html_strips_boilerplate(tmp_path):
    for name in ("a.html", "a.htm"):
        (tmp_path / name).write_text(HTML)
        t = ingest.load_text(str(tmp_path / name))
        assert "Hello world." in t and "menu" not in t and "foot" not in t and "j()" not in t


def test_unsupported(tmp_path):
    with pytest.raises(ValueError):
        ingest.load_text(str(tmp_path / "a.xyz"))


def test_docx(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("First para.")
    d.add_paragraph("Second para.")
    d.save(tmp_path / "a.docx")
    assert ingest.load_text(str(tmp_path / "a.docx")) == "First para.\nSecond para."


def test_epub(tmp_path):
    from ebooklib import epub
    book = epub.EpubBook()
    book.set_identifier("id")
    book.set_title("t")
    book.set_language("en")
    ch = epub.EpubHtml(title="c", file_name="c.xhtml", lang="en")
    ch.content = "<html><body><p>Epub body text.</p></body></html>"
    book.add_item(ch)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = [ch]
    epub.write_epub(str(tmp_path / "a.epub"), book)
    assert "Epub body text." in ingest.load_text(str(tmp_path / "a.epub"))


def test_pdf(tmp_path):
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.add_blank_page(100, 100)
    with open(tmp_path / "a.pdf", "wb") as f:
        w.write(f)
    assert ingest.load_text(str(tmp_path / "a.pdf")) == "\n"


class FakeWeb:
    """Minimal requests.get replacement keyed by URL."""

    def __init__(self, pages):
        self.pages, self.calls = pages, []

    def __call__(self, url, **kw):
        self.calls.append(url)
        if url not in self.pages:
            raise requests.HTTPError("404")
        return types.SimpleNamespace(text=self.pages[url], raise_for_status=lambda: None)


class FakeRobots:
    disallowed = ()

    def set_url(self, url):
        pass

    def read(self):
        pass

    def can_fetch(self, ua, url):
        return not url.endswith(self.disallowed) if self.disallowed else True


def test_single_url(monkeypatch):
    web = FakeWeb({"http://x.test/a": HTML})
    monkeypatch.setattr(ingest.requests, "get", web)
    assert "Hello world." in ingest.load_text("http://x.test/a")


def test_crawl(monkeypatch):
    pages = {
        "http://x.test/": '<a href="/a">a</a><a href="/a#frag">a2</a><a href="http://other.test/">o</a>'
                          '<a href="mailto:x@y">m</a><a href="/missing">m</a><a href="/private">p</a><p>root</p>',
        "http://x.test/a": '<a href="/">back</a><p>page a</p>',
        "http://x.test/private": "<p>secret</p>",
    }
    web = FakeWeb(pages)
    FakeRobots.disallowed = ("/private",)
    monkeypatch.setattr(ingest.requests, "get", web)
    monkeypatch.setattr(ingest.urllib.robotparser, "RobotFileParser", FakeRobots)
    text = ingest.load_text("http://x.test/", crawl=True, max_pages=10, delay=0)
    assert "root" in text and "page a" in text and "secret" not in text
    assert web.calls.count("http://x.test/a") == 1
    assert not any("other.test" in c for c in web.calls)


def test_crawl_max_pages_and_no_robots(monkeypatch):
    class Broken(FakeRobots):
        def read(self):
            raise OSError

    pages = {"http://x.test/": '<a href="/a">a</a><p>root</p>', "http://x.test/a": "<p>page a</p>"}
    monkeypatch.setattr(ingest.requests, "get", FakeWeb(pages))
    monkeypatch.setattr(ingest.urllib.robotparser, "RobotFileParser", Broken)
    monkeypatch.setattr(ingest.time, "sleep", lambda s: None)
    text = ingest.load_text("http://x.test/", crawl=True, max_pages=1, delay=0.5)
    assert "root" in text and "page a" not in text
